import os
import urllib.parse

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Tiến độ Tuần 3-4 trở đi"

# Enable grid lines
ws.views.sheetView[0].showGridLines = True

# Title block
ws.merge_cells("A1:E1")
title_cell = ws["A1"]
title_cell.value = "BÁO CÁO TIẾN ĐỘ THỰC HIỆN ĐỀ TÀI KHÓA LUẬN TỐT NGHIỆP"
title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F497D")
title_cell.alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells("A2:E2")
sub_cell = ws["A2"]
sub_cell.value = "Giai đoạn: Tuần 3 – Tuần 4 và Tiến độ triển khai thực nghiệm tiếp theo (21/09/2026 – Hiện tại)"
sub_cell.font = Font(name="Calibri", size=11, italic=True, color="595959")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")

# Metadata block
meta_font_bold = Font(name="Calibri", size=10, bold=True)
meta_font_val = Font(name="Calibri", size=10)

metadata = [
    ("Sinh viên thực hiện:", "Bùi Vạn Khải — MSSV: 24520719"),
    ("Giảng viên hướng dẫn:", "ThS. Trần Thị Hồng Yến"),
    ("Đơn vị đào tạo:", "Khoa Kỹ thuật Phần mềm — Trường Đại học Công nghệ Thông tin, ĐHQG-HCM"),
    ("Tên đề tài:", "Nghiên cứu, xây dựng hệ thống tự động đánh giá và phát hiện lỗi trong mã nguồn trên môi trường Pull Request")
]

for idx, (lbl, val) in enumerate(metadata, start=4):
    c_lbl = ws.cell(row=idx, column=1, value=lbl)
    c_lbl.font = meta_font_bold
    ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=5)
    c_val = ws.cell(row=idx, column=2, value=val)
    c_val.font = meta_font_val

# Table Header
headers = ["STT", "Nội dung công việc", "Trạng thái", "Đánh giá", "Minh chứng"]
header_row = 9

header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border_side = Side(border_style="thin", color="D9D9D9")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

for col_idx, h_text in enumerate(headers, start=1):
    cell = ws.cell(row=header_row, column=col_idx, value=h_text)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

ws.row_dimensions[header_row].height = 28

# GitHub Base URLs
GH_PUB_BASE = "https://github.com/Zikenic/AI-Code-Auditor-Thesis-Progress/blob/main/"
GH_W12_BASE = GH_PUB_BASE

def make_pub_link(rel_path: str) -> str:
    parts = rel_path.split("/")
    encoded = "/".join(urllib.parse.quote(p) for p in parts)
    return GH_PUB_BASE + encoded

# Rows 1 to 8 (Week 1-2 baseline items preserved from ROP_Week1-2.md)
w12_tasks = [
    (
        1,
        "Xác lập phạm vi đề tài & Kế hoạch tổng thể\n- Định hình kiến trúc 4 cấu hình thử nghiệm (A, B, C, D).\n- Phân bổ lộ trình thực hiện theo tuần.",
        "Hoàn thành",
        "Kế hoạch chi tiết, định rõ mục tiêu nghiên cứu và sự khác biệt giữa các cấu hình đánh giá thực nghiệm.",
        "benchmark_methodology.md",
        GH_W12_BASE + "methodology/benchmark_methodology.md"
    ),
    (
        2,
        "Khảo sát công cụ phân tích tĩnh (Static Analysis)\n- Đánh giá Ruff vs. Semgrep CE.\n- Đo đạc hiệu năng, độ trễ và khả năng tích hợp CI/CD.",
        "Hoàn thành",
        "Phân tích sâu 10 tiêu chí kỹ thuật; kết luận mô hình phân lớp kết hợp: Ruff quét nhanh cú pháp/Bandit, Semgrep quét taint/AST nội tệp.",
        "R1_static_analysis_ci_cd_summary.md",
        GH_W12_BASE + "research/R1_static_analysis_ci_cd_summary.md"
    ),
    (
        3,
        "Khảo sát truy xuất ngữ cảnh Repository (Repo RAG & AST)\n- Khảo sát AST Call Graph, Jedi, Tree-sitter.\n- Khảo sát Vector RAG (pgvector, text-embedding-3-small).",
        "Hoàn thành",
        "Xác lập kiến trúc ngữ cảnh hybrid: AST Graph phục vụ phân giải định nghĩa/call-site chính xác, Vector RAG phục vụ ngữ cảnh tài liệu và framework.",
        "R2_repo_context_rag_summary.md",
        GH_W12_BASE + "research/R2_repo_context_rag_summary.md"
    ),
    (
        4,
        "Nghiên cứu phân loại điểm yếu bảo mật (CWE Taxonomy)\n- Ánh xạ 6 nhóm lỗi mục tiêu sang chuẩn MITRE CWE.\n- Phân định ranh giới giữa kiểm tra đầu vào và lỗi tiêm lệnh (injection).",
        "Hoàn thành",
        "Ánh xạ chi tiết 17 mã CWE cụ thể; giải quyết triệt để ranh giới nhập nhằng giữa CWE-20 và CWE-89/22; loại bỏ CWE-285 quá trừu tượng.",
        "CWE_taxonomy_summary.md",
        GH_W12_BASE + "research/CWE_taxonomy_summary.md"
    ),
    (
        5,
        "Xây dựng đặc tả yêu cầu hệ thống (FR & NFR)\n- Soạn thảo 10 yêu cầu chức năng (FR-01 đến FR-10).\n- Soạn thảo 7 yêu cầu phi chức năng (NFR-01 đến NFR-07).",
        "Hoàn thành",
        "Đặc tả hoàn chỉnh theo mẫu kiểm thử được (testable schema), xác định rõ luồng xử lý PR, chi phí token, độ trễ CI và ngưỡng đánh giá.",
        "functional_nonfunctional_requirements.md",
        GH_W12_BASE + "requirements/functional_nonfunctional_requirements.md"
    ),
    (
        6,
        "Thiết kế bộ dữ liệu Benchmark & Lược đồ ca kiểm thử\n- Xác lập quy mô tối thiểu 100 ca (50 Buggy / 50 Clean).\n- Thiết kế lược đồ metadata 16 trường chuẩn.",
        "Hoàn thành",
        "Bộ khung benchmark chặt chẽ, cân bằng 6 nhóm lỗi, phân bổ độ phức tạp đơn tệp/đa tệp (~50/50), quy định kiểm tra hồi quy độc lập.",
        "dataset_design_test_case_list.md",
        GH_W12_BASE + "dataset/dataset_design_test_case_list.md"
    ),
    (
        7,
        "Xác lập quy tắc so khớp & Tiêu chí đo lường\n- Áp dụng dung sai ±10 dòng locus theo chuẩn Real-Vuln.\n- Xác lập quy tắc tiêu thụ 1-1 và tính lỗi FP.",
        "Hoàn thành",
        "Quy tắc so khớp khách quan, cơ chế tính điểm tự động hóa dựa trên file chuẩn hóa, CWE pre-declared và tọa độ dòng lỗi.",
        "benchmark_methodology.md",
        GH_W12_BASE + "methodology/benchmark_methodology.md"
    ),
    (
        8,
        "Khóa phương pháp luận nghiên cứu (Methodology Finalization)\n- Đồng bộ hóa toàn bộ tài liệu dự án.\n- Phân định rõ phạm vi GVHD xác nhận và phương pháp luận kỹ thuật.",
        "Hoàn thành",
        "Toàn bộ các quyết định phương pháp luận làm việc đã được chuyển trạng thái chính thức, tạo tiền đề vững chắc cho việc thu thập ca kiểm thử.",
        "benchmark_methodology.md",
        GH_W12_BASE + "methodology/benchmark_methodology.md"
    )
]

# Rows 9 to 19 (Week 3-4 and subsequent work, rigorously reconciled)
subsequent_tasks = [
    (
        9,
        "Thiết kế kiến trúc hệ thống tổng thể (System Architecture Design)",
        "Hoàn thành thiết kế; chưa triển khai",
        "Hoàn thiện thiết kế kiến trúc hệ thống 5 tầng (Ingestion, Pre-analysis, Context Retrieval, LLM Engine, Reporting & Feedback); định nghĩa luồng xử lý Webhook và CI/CD.",
        "system_architecture.md",
        make_pub_link("Project Documents/architecture/system_architecture.md")
    ),
    (
        10,
        "Thiết kế cấu trúc cơ sở dữ liệu quan hệ & Vector (Database Logical Design)",
        "Hoàn thành thiết kế; chưa triển khai",
        "Thiết kế lược đồ quan hệ logic 5 thực thể (Audits, Findings, Rules, Repositories, FileSnapshots) cùng cấu trúc AST Symbols và Code Embeddings (pgvector). Tầng bảng vật lý, Alembic migrations và production indexes chưa triển khai.",
        "system_architecture.md (Section 4)",
        make_pub_link("Project Documents/architecture/system_architecture.md")
    ),
    (
        11,
        "Thiết kế đặc tả giao diện lập trình ứng dụng & Xử lý lỗi (API Specification)",
        "Hoàn thành thiết kế; chưa triển khai",
        "Đặc tả kỹ thuật 4 REST endpoints chính (POST webhook, POST audits, GET audits/{id}, GET findings) và 1 endpoint hoãn lại (POST evaluations/run). Chuẩn hóa định dạng lỗi theo JSON error envelope nội bộ dự án (không phải RFC 7807). Xác thực Bearer token ở trạng thái đề xuất chờ duyệt. Chưa triển khai FastAPI routes.",
        "system_architecture.md (Section 7.1)",
        make_pub_link("Project Documents/architecture/system_architecture.md")
    ),
    (
        12,
        "Tích hợp công cụ phân tích tĩnh tất định (Static Analysis Integration)",
        "Hoàn thành trong phạm vi thí điểm",
        "Tích hợp engine phân tích tĩnh kết hợp Ruff và Semgrep CE; ánh xạ tập luật đóng băng (ruleset manifest 286 luật) sang finding chuẩn; hỗ trợ cơ chế nén phát hiện lặp. Hoàn thành kiểm thử đơn vị nội bộ.",
        "static_analysis/engine.py",
        make_pub_link("src/auditor/static_analysis/engine.py")
    ),
    (
        13,
        "Xây dựng bộ dữ liệu kiểm thử thí điểm (Pilot Dataset & Ground Truth)",
        "Hoàn thành trong phạm vi thí điểm",
        "Xây dựng và đóng băng bộ dữ liệu thí điểm gồm đúng 18 ca kiểm thử (6 nhóm CWE) với 23 ground-truth defects; mã băm SHA-256 bảo toàn bitwise. Mục tiêu mở rộng 100 ca kiểm thử toàn diện vẫn đang tiếp tục thu thập.",
        "data/pilot_cases.json",
        make_pub_link("data/pilot_cases.json")
    ),
    (
        14,
        "Thực nghiệm và đánh giá so sánh Cấu hình B & C (LLM & Hybrid Static-LLM)",
        "Hoàn thành trong phạm vi thí điểm",
        "Thực thi trên 18 ca thí điểm (23 defects). Config B: 100 predictions, TP 6, FP 94, FN 17, Precision 6.00%, Recall 26.09%, F1 9.76%, chi phí $1.3728. Config C (17 ca hoàn thành, CAND-0008 lỗi JSON): 95 predictions, TP 8, FP 87, FN 15, Precision 8.42%, Recall 34.78%, F1 13.56%, chi phí $1.3386. So sánh cặp (17 ca loại CAND-0008): Config B (TP 6, FP 86, FN 16, P 6.52%, R 27.27%, F1 10.53%) vs Config C (TP 8, FP 87, FN 14, P 8.42%, R 36.36%, F1 13.68%). Kết quả mang tính mô tả thực nghiệm, chưa khẳng định ý nghĩa thống kê hay quan hệ nhân quả.",
        "config_b_pilot_run.json",
        make_pub_link("data/results/config_b_pilot_run.json")
    ),
    (
        15,
        "Hiện thực hóa và thực nghiệm Cấu hình D (Repository Context Retrieval)",
        "Còn giới hạn xác minh",
        "Hiện thực hóa module trích xuất ngữ cảnh AST Graph & chunking. Thực nghiệm Config D trên 18 ca: TP 7, FP 85, FN 16, Precision 7.61%, Recall 30.43%, F1 12.17%, chi phí $1.5271. Giới hạn xác minh: ước tính ngữ cảnh từ 607–3,924 tokens dựa trên công thức heuristic floor(len/3.2)+20, chưa đo lường bằng tokenizer thực tế; trích xuất xuyên tệp chưa được xác thực toàn diện.",
        "config_d_pilot_run.json",
        make_pub_link("data/results/config_d_pilot_run.json")
    ),
    (
        16,
        "Thẩm định căn nguyên lỗi bằng sinh viên đánh giá (Phase 5E Root-Cause Review)",
        "Hoàn thành trong phạm vi thí điểm",
        "Thực hiện thẩm định căn nguyên mang tính khám phá bởi 1 sinh viên (single-student-rater) theo rubric v1.0, không phải đồng thuận chuyên gia độc lập hay thử nghiệm mù. Kết quả thẩm định trên 21 cặp so khớp cơ học TP: 15 ROOT_CAUSE_ALIGNED, 6 PROXIMITY_ONLY, 0 INDETERMINATE.",
        "adjudication_review_5e.md",
        make_pub_link("Project Documents/dataset/adjudication_review_5e.md")
    ),
    (
        17,
        "Xây dựng kịch bản kiểm diễn kho mã nguồn cục bộ (Local Repository CLI Demo)",
        "Hoàn thành",
        "Xây dựng công cụ CLI độc lập tools/demo_audit.py cho phép quét tĩnh thư mục cục bộ với các tùy chọn (--path, --format, --severity-threshold, --fail-on-findings), xuất bảng Rich/JSON và mã thoát chuẩn; đạt 7/7 unit tests. Chưa bao gồm dịch vụ web FastAPI hay tích hợp GitHub PR trực tiếp.",
        "tools/demo_audit.py",
        make_pub_link("tools/demo_audit.py")
    ),
    (
        18,
        "Kiểm thử hồi quy tự động & Bảo toàn dữ liệu (Regression & Integrity)",
        "Hoàn thành",
        "Toàn bộ bộ kiểm thử tự động toàn dự án đạt 140/140 tests passed; kiểm tra Ruff 0 lỗi; xác thực bitwise 5/5 mã băm SHA-256 của các artifact nghiên cứu đã đóng băng (pilot_cases.json và 4 tệp kết quả thực nghiệm A/B/C/D).",
        "tests/test_demo_audit.py",
        make_pub_link("tests/test_demo_audit.py")
    ),
    (
        19,
        "Rà soát và chuẩn hóa tài liệu kỹ thuật & Giấy phép (Documentation Review)",
        "Hoàn thành",
        "Chuẩn hóa Section 7.1.1 (JSON error envelope nội bộ dự án, xác thực Bearer token ở trạng thái đề xuất), đồng bộ trạng thái bàn giao AI_hand_off.md và nhật ký devlog.md; bảo vệ an toàn sở hữu trí tuệ quy tắc Semgrep trên bản phát hành công khai.",
        "devlog.md",
        make_pub_link("Project Documents/devlog.md")
    )
]

all_tasks = w12_tasks + subsequent_tasks

status_colors = {
    "Hoàn thành": {"fill": "E2EFDA", "font": "375623"},
    "Hoàn thành trong phạm vi thí điểm": {"fill": "DDEBF7", "font": "1F4E78"},
    "Hoàn thành thiết kế; chưa triển khai": {"fill": "FFF2CC", "font": "7F6000"},
    "Còn giới hạn xác minh": {"fill": "FCE4D6", "font": "C65911"},
    "Đang thực hiện": {"fill": "EDEDED", "font": "595959"}
}

row_font = Font(name="Calibri", size=10)
row_font_bold = Font(name="Calibri", size=10, bold=True)
link_font = Font(name="Calibri", size=9, color="0563C1", underline="single")

for idx, (stt, title, status, eval_text, proof_label, proof_url) in enumerate(all_tasks, start=header_row + 1):
    # STT
    c_stt = ws.cell(row=idx, column=1, value=stt)
    c_stt.font = row_font_bold
    c_stt.alignment = Alignment(horizontal="center", vertical="top")
    c_stt.border = thin_border

    # Title
    c_title = ws.cell(row=idx, column=2, value=title)
    c_title.font = row_font_bold
    c_title.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c_title.border = thin_border

    # Status
    c_status = ws.cell(row=idx, column=3, value=status)
    sc = status_colors.get(status, {"fill": "FFFFFF", "font": "000000"})
    c_status.font = Font(name="Calibri", size=10, bold=True, color=sc["font"])
    c_status.fill = PatternFill(start_color=sc["fill"], end_color=sc["fill"], fill_type="solid")
    c_status.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    c_status.border = thin_border

    # Eval
    c_eval = ws.cell(row=idx, column=4, value=eval_text)
    c_eval.font = row_font
    c_eval.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c_eval.border = thin_border

    # Proof (Hyperlink)
    c_proof = ws.cell(row=idx, column=5, value=proof_label)
    c_proof.font = link_font
    c_proof.hyperlink = proof_url
    c_proof.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c_proof.border = thin_border

    # Subtle alternating row background for non-status columns
    if idx % 2 == 1:
        alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        c_stt.fill = alt_fill
        c_title.fill = alt_fill
        c_eval.fill = alt_fill
        c_proof.fill = alt_fill

# Column widths
ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 24
ws.column_dimensions["D"].width = 54
ws.column_dimensions["E"].width = 30

# Auto-filter
ws.auto_filter.ref = f"A{header_row}:E{header_row + len(all_tasks)}"

# Freeze pane below header
ws.freeze_panes = f"A{header_row + 1}"

target_dir = r"Project Documents\reports"
os.makedirs(target_dir, exist_ok=True)
target_path = os.path.join(target_dir, "Week_3-4_and_subsequent_progress.xlsx")

wb.save(target_path)
print(f"Successfully generated: {target_path} with {len(all_tasks)} rows")
