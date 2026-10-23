"""Generate standalone Week 5-6 progress report workbook.

Creates 'Project Documents/reports/BaoCaoTienDo_DA1_Tuan5-6.xlsx' matching the
visual style of the reference report with 15 distinct tasks.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REPO_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = REPO_ROOT / "Project Documents" / "reports" / "BaoCaoTienDo_DA1_Tuan5-6.xlsx"
PUBLIC_REPO_URL = "https://github.com/Zikenic/AI-Code-Auditor-Thesis-Progress/blob/main"

TASKS = [
    {
        "stt": 1,
        "content": (
            "Hoàn thiện manifest bộ benchmark 100 ca kiểm thử\n"
            "- Chốt quy mô 100 ca: 50 ca có lỗi và 50 clean control.\n"
            "- Chuẩn hóa mã định danh và metadata cho từng ca."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Manifest ghi nhận đủ 100 ca theo tỷ lệ 50/50. Đây là mốc quy mô dữ liệu; "
            "chất lượng đại diện của từng ca vẫn cần được diễn giải cùng nguồn gốc và phạm vi kiểm thử."
        ),
        "evidence": "data/benchmark_manifest.json",
        "link_target": f"{PUBLIC_REPO_URL}/data/benchmark_manifest.json",
    },
    {
        "stt": 2,
        "content": (
            "Thu thập và materialize snapshot mã nguồn upstream\n"
            "- Lưu phiên bản mã nguồn trước khi sửa lỗi và các tệp kiểm thử liên quan.\n"
            "- Ghi nhận repository nguồn, commit, SHA-256 và vị trí lỗi."
        ),
        "status": "Hoàn thành có điều kiện",
        "eval": (
            "Báo cáo đóng gói ghi nhận 18 snapshot upstream cùng metadata provenance. "
            "Cần xem kiểm thử hash/locus như bằng chứng toàn vẹn kỹ thuật, không tự động đồng nghĩa "
            "với thẩm định độc lập về tính đại diện của benchmark."
        ),
        "evidence": "data/fixtures/upstream_snapshots/\nprovenance.json",
        "link_target": f"{PUBLIC_REPO_URL}/data/fixtures/upstream_snapshots/",
    },
    {
        "stt": 3,
        "content": (
            "Bổ sung các ca lỗi cục bộ cho benchmark\n"
            "- Tạo 32 ca lỗi tổng hợp/thiết kế thủ công để bổ sung bộ dữ liệu.\n"
            "- Liên kết từng ca với manifest và ground-truth tương ứng."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Bộ dữ liệu được báo cáo có 32 ca lỗi bổ sung ngoài các snapshot upstream. "
            "Cần giữ rõ nhãn nguồn gốc để phân biệt ca upstream với ca tự xây dựng khi phân tích kết quả."
        ),
        "evidence": "data/benchmark_manifest.json\ndata/fixtures/",
        "link_target": f"{PUBLIC_REPO_URL}/data/benchmark_manifest.json",
    },
    {
        "stt": 4,
        "content": (
            "Xây dựng bộ clean control\n"
            "- Chuẩn hóa 50 ca đối chứng sạch.\n"
            "- Gồm 18 ca thủ công và 32 ca tổng hợp theo báo cáo nghiệm thu."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Manifest phân tách 50 clean control khỏi 50 ca lỗi, tạo cơ sở đánh giá false positive. "
            "Khi báo cáo kết quả, cần giữ riêng kết quả trên nhóm đối chứng và không xem nhãn clean "
            "là bằng chứng tuyệt đối nếu chưa có xác minh độc lập."
        ),
        "evidence": "data/benchmark_manifest.json",
        "link_target": f"{PUBLIC_REPO_URL}/data/benchmark_manifest.json",
    },
    {
        "stt": 5,
        "content": (
            "Viết kiểm thử toàn vẹn và cấu trúc benchmark\n"
            "- Kiểm tra số lượng, phân bố nhãn và schema manifest.\n"
            "- Kiểm tra SHA-256 và vị trí dòng lỗi của snapshot."
        ),
        "status": "Hoàn thành có điều kiện",
        "eval": (
            "Có bộ kiểm thử chuyên biệt cho schema, cân bằng 50/50, checksum và locus. "
            "Phạm vi test chạy ở private checkout và public snapshot không giống nhau; "
            "báo cáo kết quả phải ghi riêng từng phạm vi."
        ),
        "evidence": "tests/test_benchmark_100_provenance.py",
        "link_target": f"{PUBLIC_REPO_URL}/tests/test_benchmark_100_provenance.py",
    },
    {
        "stt": 6,
        "content": (
            "Tích hợp cơ chế đếm token tương thích model thực tế\n"
            "- Gọi endpoint đếm token native qua gateway Anthropic.\n"
            "- Cấu hình model identifier sử dụng cho phép đo."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Module sử dụng endpoint POST /v1/messages/count_tokens với model được cấu hình. "
            "Phép đo phản ánh cấu trúc request của endpoint được gọi, không nên mặc định "
            "suy rộng sang mọi tokenizer/model khác."
        ),
        "evidence": "src/auditor/llm/token_counter.py\ntests/test_token_counting.py",
        "link_target": f"{PUBLIC_REPO_URL}/src/auditor/llm/token_counter.py",
    },
    {
        "stt": 7,
        "content": (
            "Hạch toán prompt và giới hạn ngân sách ngữ cảnh\n"
            "- Tính token trên request đã render, gồm system/user content và cấu trúc liên quan được gửi.\n"
            "- Dự phòng 16.384 token cho completion và chặn request vượt giới hạn cấu hình."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Báo cáo nghiệm thu ghi nhận giới hạn context 200.000 token và completion reserve "
            "16.384 token. Hành vi phải được đánh giá theo model/gateway thực sự cấu hình, "
            "không xem các con số này là giới hạn chung cho mọi model."
        ),
        "evidence": "src/auditor/llm/token_counter.py\ntests/test_token_counting.py",
        "link_target": f"{PUBLIC_REPO_URL}/src/auditor/llm/token_counter.py",
    },
    {
        "stt": 8,
        "content": (
            "Kiểm thử lỗi endpoint và cơ chế fallback đếm token\n"
            "- Kiểm thử xử lý lỗi khi endpoint native không khả dụng.\n"
            "- Tách biệt ước lượng ký tự và TikToken BPE khỏi phép đếm native; phát cảnh báo khi dùng fallback."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Các fallback được mô tả là xấp xỉ và cần giữ cảnh báo telemetry rõ ràng. "
            "Không được báo cáo kết quả fallback như phép đo token chính xác của tokenizer đích."
        ),
        "evidence": "tests/test_token_counting.py",
        "link_target": f"{PUBLIC_REPO_URL}/tests/test_token_counting.py",
    },
    {
        "stt": 9,
        "content": (
            "Triển khai backend API bằng FastAPI\n"
            "- Tổ chức các thành phần API theo module/router.\n"
            "- Kết nối luồng tiếp nhận sự kiện và dữ liệu audit của hệ thống."
        ),
        "status": "Hoàn thành trong phạm vi prototype",
        "eval": (
            "Backend FastAPI đã được triển khai và có kiểm thử. Đây là mốc prototype; "
            "không đồng nghĩa đã triển khai/kiểm chứng trong môi trường production "
            "hoặc hoàn tất tích hợp GitHub Pull Request đầu-cuối."
        ),
        "evidence": "src/auditor/api/",
        "link_target": f"{PUBLIC_REPO_URL}/src/auditor/api/",
    },
    {
        "stt": 10,
        "content": (
            "Thiết lập lưu trữ CSDL vật lý bằng SQLite WAL\n"
            "- Lưu dữ liệu nghiệp vụ thay vì chỉ dùng trạng thái tạm trong bộ nhớ.\n"
            "- Kiểm tra schema và khả năng duy trì dữ liệu qua lần chạy."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Báo cáo nghiệm thu ghi nhận file SQLite vật lý với WAL và các bảng repositories, "
            "audits, findings. Bằng chứng persistence cần được đọc trong phạm vi test hiện có."
        ),
        "evidence": "tests/test_api_physical_db.py\ntests/test_db_persistence.py",
        "link_target": f"{PUBLIC_REPO_URL}/tests/test_api_physical_db.py",
    },
    {
        "stt": 11,
        "content": (
            "Xác thực chữ ký webhook GitHub\n"
            "- Kiểm tra chữ ký HMAC-SHA256 từ header X-Hub-Signature-256.\n"
            "- Từ chối payload có chữ ký không hợp lệ hoặc bị thay đổi."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Triển khai xác thực chữ ký với so sánh constant-time qua hmac.compare_digest "
            "theo báo cáo kỹ thuật. Cần cấu hình secret an toàn khi triển khai thực tế."
        ),
        "evidence": "src/auditor/api/webhooks.py\ntests/test_webhook_intake.py",
        "link_target": f"{PUBLIC_REPO_URL}/src/auditor/api/webhooks.py",
    },
    {
        "stt": 12,
        "content": (
            "Kiểm soát webhook lặp và phạm vi workspace\n"
            "- Dùng delivery ID để nhận diện sự kiện đã xử lý.\n"
            "- Giới hạn đường dẫn repository trong allowlist để ngăn truy cập vượt phạm vi."
        ),
        "status": "Hoàn thành có điều kiện",
        "eval": (
            "Có kiểm thử deduplication và kiểm tra đường dẫn. Việc kích hoạt xử lý qua "
            "FastAPI BackgroundTasks phù hợp với prototype nhưng không thay thế hàng đợi "
            "bền vững có retry/quan sát cho tải production."
        ),
        "evidence": "src/auditor/api/webhooks.py\ntests/test_webhook_intake.py",
        "link_target": f"{PUBLIC_REPO_URL}/src/auditor/api/webhooks.py",
    },
    {
        "stt": 13,
        "content": (
            "Chạy kiểm thử hồi quy, kiểm thử ablation và lint\n"
            "- Chạy bộ test của private checkout và bộ ablation riêng.\n"
            "- Kiểm tra Ruff và ghi nhận riêng kết quả kiểm thử ở public snapshot."
        ),
        "status": "Hoàn thành có điều kiện",
        "eval": (
            "Báo cáo tác vụ ghi nhận 306/306 test trong private checkout; public staging báo cáo "
            "43 passed, 1 skipped. Hai con số thuộc phạm vi chạy khác nhau và không được gộp "
            "thành một kết quả test duy nhất."
        ),
        "evidence": "tests/\ndata/fixtures/ablation_fixtures_v2/\npyproject.toml",
        "link_target": f"{PUBLIC_REPO_URL}/tests/",
    },
    {
        "stt": 14,
        "content": (
            "Đồng bộ tài liệu kiến trúc và bàn giao sau mốc Tuần 5–6\n"
            "- Cập nhật AI hand-off, devlog và manifesto theo hiện trạng triển khai.\n"
            "- Ghi nhận giới hạn còn tồn tại và ưu tiên Tuần 7–8."
        ),
        "status": "Hoàn thành",
        "eval": (
            "Tài liệu bàn giao đã được cập nhật theo báo cáo. Các mục còn là prototype hoặc "
            "mới được kiểm thử trong phạm vi nội bộ cần giữ nguyên chú thích giới hạn, "
            "không chuyển thành tuyên bố production-ready."
        ),
        "evidence": "Project Documents/AI_hand_off.md\nProject Documents/devlog.md\nProject Documents/manifesto.md",
        "link_target": f"{PUBLIC_REPO_URL}/Project%20Documents/AI_hand_off.md",
    },
    {
        "stt": 15,
        "content": (
            "Đóng gói báo cáo nghiệm thu và công bố bản repository đã làm sạch\n"
            "- Đối soát bằng chứng, workbook và các tài liệu mốc.\n"
            "- Loại trừ artefact nội bộ/restricted trước khi cập nhật public repository."
        ),
        "status": "Hoàn thành có điều kiện",
        "eval": (
            "Báo cáo ghi nhận public commit 6f31a7f và loại trừ các tệp/rule pack nội bộ đã nêu. "
            "Cần tiếp tục xác minh các liên kết công khai và rà soát toàn bộ diff trước mọi lần công bố tiếp theo."
        ),
        "evidence": "reports/week5_6_deliverables_acceptance_report.md\n.gitignore",
        "link_target": f"{PUBLIC_REPO_URL}/reports/week5_6_deliverables_acceptance_report.md",
    },
]


def create_workbook() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tiến độ Tuần 5-6"

    # Fonts
    title_font = Font(name="Calibri", size=14, bold=True, color="000000")
    subtitle_font = Font(name="Calibri", size=11, bold=False, italic=True, color="000000")
    label_font = Font(name="Calibri", size=10, bold=True, color="000000")
    value_font = Font(name="Calibri", size=10, bold=False, color="000000")

    ws["A1"] = "BÁO CÁO TIẾN ĐỘ THỰC HIỆN ĐỀ TÀI KHÓA LUẬN TỐT NGHIỆP"
    ws["A1"].font = title_font

    ws["A2"] = "Giai đoạn: Tuần 5 – Tuần 6 (05/10/2026 – 18/10/2026)"
    ws["A2"].font = subtitle_font

    metadata = [
        ("A4", "Sinh viên thực hiện:", "B4", "Bùi Vạn Khải — MSSV: 24520719"),
        ("A5", "Giảng viên hướng dẫn:", "B5", "ThS. Trần Thị Hồng Yến"),
        ("A6", "Đơn vị đào tạo:", "B6", "Khoa Kỹ thuật Phần mềm — Trường Đại học Công nghệ Thông tin, ĐHQG-HCM"),
        (
            "A7",
            "Tên đề tài:",
            "B7",
            "Nghiên cứu, xây dựng hệ thống tự động đánh giá và phát hiện lỗi trong mã nguồn trên môi trường Pull Request",
        ),
    ]

    for a_cell, a_val, b_cell, b_val in metadata:
        ws[a_cell] = a_val
        ws[a_cell].font = label_font
        ws[b_cell] = b_val
        ws[b_cell].font = value_font

    # Column dimensions
    ws.column_dimensions["A"].width = 8.0
    ws.column_dimensions["B"].width = 38.0
    ws.column_dimensions["C"].width = 24.0
    ws.column_dimensions["D"].width = 54.0
    ws.column_dimensions["E"].width = 32.0

    # Header row
    header_fill = PatternFill(start_color="001F497D", end_color="001F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="00FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="00D9D9D9"),
        right=Side(style="thin", color="00D9D9D9"),
        top=Side(style="thin", color="00D9D9D9"),
        bottom=Side(style="thin", color="00D9D9D9"),
    )

    headers = ["STT", "Nội dung công việc", "Trạng thái", "Đánh giá", "Minh chứng"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=9, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[9].height = 25.0

    # Task rows
    body_font = Font(name="Calibri", size=10, bold=False, color="000000")
    stt_font = Font(name="Calibri", size=10, bold=True, color="000000")
    link_font = Font(name="Calibri", size=10, bold=False, color="000563C1", underline="single")

    status_colors = {
        "Hoàn thành": "00375623",
        "Hoàn thành có điều kiện": "007F6000",
        "Hoàn thành trong phạm vi prototype": "001F4E78",
    }

    start_row = 10
    for idx, t in enumerate(TASKS):
        row = start_row + idx

        # STT
        c1 = ws.cell(row=row, column=1, value=t["stt"])
        c1.font = stt_font
        c1.alignment = Alignment(horizontal="center", vertical="top")
        c1.border = thin_border

        # Nội dung công việc
        c2 = ws.cell(row=row, column=2, value=t["content"])
        c2.font = body_font
        c2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c2.border = thin_border

        # Trạng thái
        status_color = status_colors.get(t["status"], "000000")
        status_font = Font(name="Calibri", size=10, bold=True, color=status_color)
        c3 = ws.cell(row=row, column=3, value=t["status"])
        c3.font = status_font
        c3.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        c3.border = thin_border

        # Đánh giá
        c4 = ws.cell(row=row, column=4, value=t["eval"])
        c4.font = body_font
        c4.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c4.border = thin_border

        # Minh chứng
        c5 = ws.cell(row=row, column=5, value=t["evidence"])
        c5.font = link_font
        c5.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c5.border = thin_border
        if t.get("link_target"):
            c5.hyperlink = t["link_target"]

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH} with {len(TASKS)} tasks successfully.")


if __name__ == "__main__":
    create_workbook()
